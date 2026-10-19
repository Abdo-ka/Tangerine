"""End-to-end local worker tests. Temporary fixtures; no user files or network."""

import hashlib
import importlib.util
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import pymupdf as fitz
from PIL import Image

spec = importlib.util.spec_from_file_location(
    "engine", Path(__file__).resolve().parents[2] / "Resources/engine.py"
)
e = importlib.util.module_from_spec(spec)
spec.loader.exec_module(e)


class EngineTests(unittest.TestCase):
    def test_pages_editions_pass_paths_as_arguments(self):
        source = self.root / 'a "quoted" input.pages'
        output = self.dest / "output.docx"
        for bundle_id in ("com.apple.iWork.Pages", "com.apple.Pages"):
            for operation in (e.pages_import, e.pages_export):
                with (
                    patch.object(e, "pages_application_id", return_value=bundle_id),
                    patch.object(e, "run") as run,
                ):
                    operation(source, output)
                    args = run.call_args.args[0]
                    self.assertIn(f'tell application id "{bundle_id}"', args[2])
                    self.assertNotIn(str(source), args[2])
                    self.assertEqual(args[-2:], [source, output])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "inputs"
        self.source.mkdir()
        self.dest = self.root / "outputs"
        self.dest.mkdir()
        self.image = self.source / "photo.png"
        Image.new("RGB", (240, 160), "#2f7055").save(self.image)
        self.other = self.source / "second.png"
        Image.new("RGB", (160, 240), "#fd763d").save(self.other)
        self.pdf = self.source / "notes.pdf"
        d = fitz.open()
        for text in [
            "First page: local conversion.",
            "Second page: preserve originals.",
        ]:
            page = d.new_page()
            page.insert_text((50, 60), text)
        d.save(self.pdf)

    def tearDown(self):
        self.temp.cleanup()

    def execute(self, src, operation="convert", format="", options=None):
        inputs = src if isinstance(src, list) else [src]
        original = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs if p.is_file()}
        result = e.execute(
            dict(
                inputs=[str(p) for p in inputs],
                destination=str(self.dest),
                operation=operation,
                format=format,
                options=options or {},
            )
        )
        for p, checksum in original.items():
            self.assertEqual(checksum, hashlib.sha256(p.read_bytes()).hexdigest())
        paths = [Path(p) for p in result["outputs"]]
        for p in paths:
            self.assertTrue(p.exists())
            self.assertTrue(any(p.iterdir()) if p.is_dir() else p.stat().st_size > 0)
        return paths

    def test_image_formats(self):
        for fmt in ["jpg", "png", "webp", "heic", "tiff", "avif", "bmp", "pdf"]:
            with self.subTest(format=fmt):
                output = self.execute(self.image, format=fmt)[0]
                if fmt == "pdf":
                    self.assertEqual(len(fitz.open(output)), 1)
                else:
                    with Image.open(output) as image:
                        self.assertEqual(image.size, (240, 160))

    def test_all_image_tools(self):
        tools = {
            "compress": {"scale": 0.5, "quality": 60},
            "removeMetadata": {},
            "editImage": {"brightness": 1.3, "rotation": 90},
            "annotateImage": {
                "annotations": [
                    {"type": "arrow", "x": 0.1, "y": 0.1, "width": 0.5, "height": 0.5},
                    {"type": "text", "text": "Hello", "x": 0.1, "y": 0.1},
                ]
            },
            "frameImage": {"spacing": 20, "radius": 10, "gradient": True},
            "cropImage": {"x": 0.25, "y": 0, "width": 0.5, "height": 1},
            "redactImage": {
                "x": 0.1,
                "y": 0.1,
                "width": 0.5,
                "height": 0.5,
                "redaction": "blur",
            },
            "createPDF": {},
            "createCollage": {"layout": "featured", "cellSize": 100},
        }
        for tool, options in tools.items():
            with self.subTest(tool=tool):
                paths = self.execute(
                    [self.image, self.other]
                    if tool in {"createPDF", "createCollage"}
                    else self.image,
                    tool,
                    options=options,
                )
                if tool == "cropImage":
                    with Image.open(paths[0]) as image:
                        self.assertEqual(image.size, (120, 160))
                if tool == "createPDF":
                    self.assertEqual(len(fitz.open(paths[0])), 2)

    def test_svg(self):
        svg = self.source / "icon.svg"
        svg.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><rect width="100" height="100" fill="red"/></svg>'
        )
        with Image.open(self.execute(svg, format="png")[0]) as image:
            self.assertEqual(image.size, (200, 200))

    def test_pdf_tools_and_exports(self):
        for tool, options in [
            ("compress", {}),
            ("removeMetadata", {"metadata": {"title": "Local"}}),
            ("splitPDF", {}),
            ("mergePDF", {}),
            ("organizePDF", {"pages": "2,1,1", "rotation": 90}),
        ]:
            with self.subTest(tool=tool):
                output = self.execute(
                    [self.pdf, self.pdf] if tool == "mergePDF" else self.pdf,
                    tool,
                    options=options,
                )[0]
                if tool == "splitPDF":
                    self.assertEqual(len(list(output.glob("*.pdf"))), 2)
                if tool == "organizePDF":
                    self.assertEqual(len(fitz.open(output)), 3)
                    self.assertEqual(fitz.open(output)[0].rotation, 90)
        for fmt in ["txt", "md", "jpg", "png", "docx", "pptx", "xlsx", "epub"]:
            with self.subTest(format=fmt):
                self.execute(self.pdf, format=fmt)
        txt = self.execute(self.pdf, format="txt")[0].read_text()
        self.assertIn("First page", txt)
        self.assertIn("Second page", txt)

    def test_document_conversions(self):
        text = self.source / "readme.md"
        text.write_text("# Hello\n\nLocal files stay local.\n\nUnicode: مرحبا — café")
        for fmt in ["txt", "docx", "pdf", "html", "epub", "png", "jpg"]:
            with self.subTest(format=fmt):
                self.execute(text, format=fmt)
        docx = self.execute(text, format="docx")[0]
        self.assertIn("Local files", self.execute(docx, format="txt")[0].read_text())
        epub = self.execute(text, format="epub")[0]
        self.assertIn("Local files", self.execute(epub, format="txt")[0].read_text())

    def test_subtitles_preserve_timing(self):
        src = self.source / "captions.srt"
        src.write_text(
            "1\n00:00:01,250 --> 00:00:03,500\nHello world\n\n2\n00:00:04,000 --> 00:00:05,000\nSecond cue\n"
        )
        vtt = self.execute(src, format="vtt")[0]
        self.assertIn("00:00:01.250", vtt.read_text())
        srt = self.execute(vtt, format="srt")[0]
        self.assertIn("00:00:03,500", srt.read_text())
        self.assertEqual(
            self.execute(src, format="txt")[0].read_text(), "Hello world\nSecond cue\n"
        )

    def test_archives_roundtrip(self):
        src = self.source / "project.zip"
        with zipfile.ZipFile(src, "w") as z:
            z.writestr("folder/hello.txt", "Original content")
        for fmt in ["zip", "tar", "gz"]:
            with self.subTest(format=fmt):
                archive = self.execute(src, format=fmt)[0]
                folder = self.execute(archive, "extractArchive")[0]
                self.assertEqual((folder / "folder/hello.txt").read_text(), "Original content")

    def test_archive_traversal_and_links_rejected(self):
        for name, mode in [
            ("../escaped.txt", 0),
            ("/absolute.txt", 0),
            ("link", 0o120777),
        ]:
            with self.subTest(name=name):
                src = self.source / "malicious.zip"
                with zipfile.ZipFile(src, "w") as z:
                    info = zipfile.ZipInfo(name)
                    info.external_attr = mode << 16
                    z.writestr(info, "outside")
                with self.assertRaises(e.ConversionError):
                    self.execute(src, "extractArchive")
        self.assertFalse((self.root / "escaped.txt").exists())
        self.assertEqual(list(self.dest.iterdir()), [])

    def test_collision_and_batch(self):
        paths = self.execute([self.image, self.other], format="jpg")
        self.assertEqual(len(paths), 2)
        again = self.execute(self.image, format="jpg")[0]
        self.assertNotEqual(paths[0], again)
        self.assertTrue(paths[0].exists())

    def test_batch_keeps_successes_and_reports_failures(self):
        invalid = self.source / "broken.png"
        invalid.write_bytes(b"not an image")
        result = e.execute(
            dict(
                inputs=[str(self.image), str(invalid)],
                destination=str(self.dest),
                format="jpg",
            )
        )
        self.assertEqual(len(result["outputs"]), 1)
        self.assertEqual(len(result["errors"]), 1)
        self.assertIn("broken.png", result["errors"][0]["input"])
        self.assertTrue(Path(result["outputs"][0]).exists())
        self.assertEqual(len(list(self.dest.iterdir())), 1)

    def test_fixed_layout_epub_and_page_rotations(self):
        output = self.execute(self.pdf, format="epub", options={"epubLayout": "fixed"})[0]
        with zipfile.ZipFile(output) as archive:
            self.assertEqual(len([n for n in archive.namelist() if n.endswith(".png")]), 2)
            package = next(n for n in archive.namelist() if n.endswith(".opf"))
            self.assertIn(b"pre-paginated", archive.read(package))
        organized = self.execute(
            self.pdf,
            "organizePDF",
            options={
                "pageSelections": [
                    {"page": 2, "rotation": 90},
                    {"page": 1, "rotation": 180},
                ]
            },
        )[0]
        doc = fitz.open(organized)
        self.assertEqual(doc[0].rotation, 90)
        self.assertEqual(doc[1].rotation, 180)

    def test_png_metadata_can_be_edited_and_removed(self):
        edited = self.execute(
            self.image,
            "removeMetadata",
            options={"metadata": {"author": "Fixture author"}},
        )[0]
        with Image.open(edited) as image:
            self.assertEqual(image.info["author"], "Fixture author")
        clean = self.execute(edited, "removeMetadata")[0]
        with Image.open(clean) as image:
            self.assertNotIn("author", image.info)

    def test_rar_stream_reader_enforces_paths(self):
        # The same libarchive streaming reader handles ZIP too, allowing a generated
        # fixture to verify its traversal checks without a proprietary RAR writer.
        valid = self.source / "stream.zip"
        with zipfile.ZipFile(valid, "w") as archive:
            archive.writestr("nested/file.txt", "content")
        folder = self.dest / "stream"
        folder.mkdir()
        e.extract_rar(valid, folder, 10000)
        self.assertEqual((folder / "nested/file.txt").read_text(), "content")
        invalid = self.source / "unsafe.zip"
        with zipfile.ZipFile(invalid, "w") as archive:
            archive.writestr("../escape.txt", "bad")
        with self.assertRaises(e.ConversionError):
            e.extract_rar(invalid, folder, 10000)

    def test_clipping_resource_fork(self):
        if __import__("sys").platform != "darwin":
            self.skipTest("macOS resource forks")
        import struct

        text = "Finder clipping text"
        data = text.encode("utf-16-be")
        data_offset = 256
        map_offset = data_offset + 4 + len(data)
        resource = bytearray(map_offset + 50)
        struct.pack_into(">4I", resource, 0, data_offset, map_offset, 4 + len(data), 50)
        struct.pack_into(">I", resource, data_offset, len(data))
        resource[data_offset + 4 : map_offset] = data
        struct.pack_into(">H", resource, map_offset + 24, 28)
        types = map_offset + 28
        struct.pack_into(">H", resource, types, 0)
        struct.pack_into(">4sHH", resource, types + 2, b"utxt", 0, 10)
        # One resource reference; 24-bit data offset begins at byte 5 and is zero.
        source = self.source / "note.textClipping"
        source.write_bytes(b"")
        Path(str(source) + "/..namedfork/rsrc").write_bytes(bytes(resource))
        self.assertEqual(self.execute(source, format="txt")[0].read_text(), text)

    def test_resize_preserves_aspect_with_one_dimension(self):
        output = self.execute(self.image, "resizeImage", options={"outputHeight": 80})[0]
        with Image.open(output) as image:
            self.assertEqual(image.size, (120, 80))
        output = self.execute(
            self.image, "resizeImage", options={"outputWidth": 90, "outputHeight": 90}
        )[0]
        with Image.open(output) as image:
            self.assertEqual(image.size, (90, 90))

    def test_cancel_exits_batch_instead_of_continuing(self):
        from unittest.mock import patch

        # SystemExit must escape per-file error aggregation and unwind staging.
        with patch.object(e, "image_tool", side_effect=SystemExit(130)) as operation:
            with self.assertRaises(SystemExit):
                e.execute(
                    dict(
                        inputs=[str(self.image), str(self.other)],
                        destination=str(self.dest),
                        format="jpg",
                    )
                )
            self.assertEqual(operation.call_count, 1)
        self.assertEqual(list(self.dest.iterdir()), [])

    def test_invalid_crop_is_atomic(self):
        with self.assertRaises(e.ConversionError):
            self.execute(self.image, "cropImage", options={"x": 0.8, "width": 0.5})
        self.assertEqual(list(self.dest.iterdir()), [])

    def test_spreadsheet_formula_injection_is_escaped(self):
        self.assertEqual(e.safe_cell('=HYPERLINK("bad")'), '\'=HYPERLINK("bad")')
        self.assertEqual(e.safe_cell("normal"), "normal")


def office_available():
    try:
        return bool(e.executable("soffice"))
    except e.ConversionError:
        return False


@unittest.skipUnless(office_available(), "LibreOffice is required for office integration tests")
class OfficeTests(unittest.TestCase):
    def test_spreadsheet_and_slides_roundtrip(self):
        from openpyxl import load_workbook
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "budget.csv"
            source.write_text("Name,Amount\nApple,12\nOrange,24\n")

            def convert(src, fmt):
                return Path(e.execute(dict(inputs=[str(src)], format=fmt))["outputs"][0])

            xlsx = convert(source, "xlsx")
            workbook = load_workbook(xlsx)
            self.assertEqual(workbook.active["A2"].value, "Apple")
            self.assertEqual(workbook.active["B3"].value, 24)
            for fmt in ["ods", "csv", "tsv", "pdf"]:
                with self.subTest(format=fmt):
                    output = convert(xlsx, fmt)
                    self.assertGreater(output.stat().st_size, 0)
                    if fmt == "tsv":
                        self.assertIn("Apple\t12", output.read_text())
                    if fmt == "pdf":
                        self.assertIn("Orange", fitz.open(output)[0].get_text())
            pptx = root / "slides.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = "Local slides"
            slide.placeholders[1].text = "Fixture content"
            deck.save(pptx)
            pdf = convert(pptx, "pdf")
            self.assertIn("Local slides", fitz.open(pdf)[0].get_text())
            images = convert(pptx, "png")
            self.assertEqual(len(list(images.glob("*.png"))), 1)
            docx = root / "document.docx"
            from docx import Document

            doc = Document()
            doc.add_heading("Layout fixture", 0)
            doc.add_paragraph("Preserve the text")
            doc.save(docx)
            self.assertIn("Preserve the text", fitz.open(convert(docx, "pdf"))[0].get_text())


@unittest.skipUnless(shutil.which("ffmpeg"), "FFmpeg is required for media integration tests")
class MediaTests(EngineTests):
    # Share fixture helpers without repeating inherited test cases.
    test_image_formats = None
    test_all_image_tools = None
    test_svg = None
    test_pdf_tools_and_exports = None
    test_document_conversions = None
    test_subtitles_preserve_timing = None
    test_archives_roundtrip = None
    test_archive_traversal_and_links_rejected = None
    test_collision_and_batch = None
    test_batch_keeps_successes_and_reports_failures = None
    test_fixed_layout_epub_and_page_rotations = None
    test_png_metadata_can_be_edited_and_removed = None
    test_rar_stream_reader_enforces_paths = None
    test_clipping_resource_fork = None
    test_resize_preserves_aspect_with_one_dimension = None
    test_cancel_exits_batch_instead_of_continuing = None
    test_invalid_crop_is_atomic = None
    test_spreadsheet_formula_injection_is_escaped = None

    def test_pages_editions_pass_paths_as_arguments(self):
        source = self.root / 'a "quoted" input.pages'
        output = self.dest / "output.docx"
        for bundle_id in ("com.apple.iWork.Pages", "com.apple.Pages"):
            for operation in (e.pages_import, e.pages_export):
                with (
                    patch.object(e, "pages_application_id", return_value=bundle_id),
                    patch.object(e, "run") as run,
                ):
                    operation(source, output)
                    args = run.call_args.args[0]
                    self.assertIn(f'tell application id "{bundle_id}"', args[2])
                    self.assertNotIn(str(source), args[2])
                    self.assertEqual(args[-2:], [source, output])

    def setUp(self):
        super().setUp()
        self.video = self.source / "clip.mp4"
        self.audio = self.source / "tone.wav"
        e.run(
            [
                e.executable("ffmpeg"),
                "-hide_banner",
                "-loglevel",
                "error",
                "-f",
                "lavfi",
                "-i",
                "testsrc2=size=160x120:rate=15",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:sample_rate=48000",
                "-t",
                "2",
                "-c:v",
                "libx264",
                "-pix_fmt",
                "yuv420p",
                "-c:a",
                "aac",
                self.video,
            ]
        )
        e.run(
            [
                e.executable("ffmpeg"),
                "-hide_banner",
                "-loglevel",
                "error",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:sample_rate=48000",
                "-t",
                "2",
                self.audio,
            ]
        )

    def test_video_formats(self):
        for fmt in ["mov", "mkv", "webm", "avi", "wmv", "gif", "mp3"]:
            with self.subTest(format=fmt):
                self.execute(self.video, format=fmt)

    def test_audio_formats(self):
        for fmt in ["mp3", "m4a", "flac", "ogg", "opus", "aiff", "wma"]:
            with self.subTest(format=fmt):
                self.execute(self.audio, format=fmt)

    def test_ogg_without_vorbis_preserves_audio_using_opus(self):
        with patch.object(e, "media_encoders", return_value={"libopus"}):
            output = self.execute(self.audio, format="ogg")[0]
        info = e.probe(output)
        self.assertEqual(info["streams"][0]["codec_name"], "opus")
        self.assertAlmostEqual(float(info["format"]["duration"]), 2, delta=0.1)

    def test_transparent_animation_to_webm(self):
        source = self.root / "animation.gif"
        frames = [Image.new("RGBA", (88, 64), color) for color in ("red", "blue")]
        for frame in frames:
            frame.putpixel((0, 0), (0, 0, 0, 0))
        frames[0].save(source, save_all=True, append_images=frames[1:], duration=150, loop=0)
        output = self.execute(source, format="webm")[0]
        info = e.probe(output)
        video = next(stream for stream in info["streams"] if stream["codec_type"] == "video")
        self.assertEqual(video["codec_name"], "vp9")
        self.assertEqual((video["width"], video["height"]), (88, 64))
        self.assertAlmostEqual(float(info["format"]["duration"]), 0.3, delta=0.1)

    def test_all_video_tools(self):
        tools = {
            "compress": {},
            "removeMetadata": {},
            "muteVideo": {},
            "trimVideo": {"start": 0.2, "end": 1.2},
            "splitVideo": {"times": ".5,1"},
            "cropVideo": {"x": 0, "y": 0, "width": 0.5, "height": 0.5},
            "changeVideoSpeed": {"speed": 2},
            "joinVideos": {},
            "videoSnapshots": {"times": ".5,1"},
            "redactVideo": {
                "x": 0.1,
                "y": 0.1,
                "width": 0.5,
                "height": 0.5,
                "start": 0.2,
                "end": 1,
                "redaction": "blur",
            },
        }
        for tool, options in tools.items():
            with self.subTest(tool=tool):
                output = self.execute(
                    [self.video, self.video] if tool == "joinVideos" else self.video,
                    tool,
                    options=options,
                )[0]
                if tool == "muteVideo":
                    self.assertFalse(
                        any(s["codec_type"] == "audio" for s in e.probe(output)["streams"])
                    )
                if tool == "trimVideo":
                    self.assertAlmostEqual(
                        float(e.probe(output)["format"]["duration"]), 1, delta=0.15
                    )
                if tool == "cropVideo":
                    self.assertEqual(e.probe(output)["streams"][0]["width"], 80)
                if tool == "changeVideoSpeed":
                    self.assertAlmostEqual(
                        float(e.probe(output)["format"]["duration"]), 1, delta=0.2
                    )
                if tool == "splitVideo":
                    self.assertEqual(len(list(output.glob("*.mp4"))), 3)
                if tool == "videoSnapshots":
                    self.assertEqual(len(list(output.glob("*.png"))), 2)

    def test_all_audio_tools(self):
        tools = {
            "compress": {},
            "removeMetadata": {},
            "normalizeAudio": {},
            "trimAudio": {"start": 0.2, "end": 1.5},
            "audioChannels": {"channels": 2},
            "audioToVideo": {"outputWidth": 320, "outputHeight": 180},
            "redactAudio": {"ranges": ".2:.4,.8:1"},
        }
        for tool, options in tools.items():
            with self.subTest(tool=tool):
                output = self.execute(self.audio, tool, options=options)[0]
                if tool == "audioChannels":
                    self.assertEqual(e.probe(output)["streams"][0]["channels"], 2)
                if tool == "audioToVideo":
                    self.assertTrue(
                        any(s["codec_type"] == "video" for s in e.probe(output)["streams"])
                    )

    def test_invalid_trim_reports_error(self):
        with self.assertRaises(e.ConversionError):
            self.execute(self.video, "trimVideo", options={"start": 3, "end": 2})


if __name__ == "__main__":
    unittest.main()

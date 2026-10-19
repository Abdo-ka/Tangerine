#!/usr/bin/env python3
"""Local conversion worker. One JSON request per process; originals are never overwritten."""

import ctypes
import ctypes.util
import errno
import html
import io
import json
import math
import os
import plistlib
import re
import shutil
import signal
import subprocess
import sys
import tarfile
import tempfile
import zipfile
from functools import lru_cache
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps
from pillow_heif import register_heif_opener

register_heif_opener()
IMAGE_EXTS = {
    "jpg",
    "jpeg",
    "png",
    "webp",
    "heic",
    "heif",
    "tif",
    "tiff",
    "avif",
    "bmp",
    "svg",
}
AUDIO_EXTS = {"mp3", "m4a", "wav", "flac", "ogg", "opus", "aiff", "aif", "wma"}
VIDEO_EXTS = {"mp4", "mov", "mkv", "webm", "avi", "wmv", "gif"}
ARCHIVE_EXTS = {"zip", "tar", "gz", "tgz", "rar"}


class ConversionError(Exception):
    pass


def executable(name):
    root = (
        Path(sys.executable).parent
        if getattr(sys, "frozen", False)
        else Path(__file__).resolve().parent
    )
    bundled = root / "bin" / name
    candidates = [
        os.environ.get("TANGERINE_" + name.upper()),
        str(bundled),
        shutil.which(name),
        "/opt/homebrew/bin/" + name,
        "/usr/local/bin/" + name,
    ]
    if name == "soffice":
        candidates[0:0] = [
            str(root / "LibreOffice.app/Contents/MacOS/soffice"),
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file() and os.access(candidate, os.X_OK):
            return candidate
    raise ConversionError(f"{name} is unavailable. Run scripts/bootstrap.sh and rebuild the app.")


_active_process = None


def cancel_worker(signum, frame):
    if _active_process is not None:
        try:
            os.killpg(_active_process.pid, signal.SIGTERM)
        except ProcessLookupError:
            pass
        _active_process.wait()
    raise SystemExit(130)


def run(args):
    """No shell: filenames and user-entered values are literal arguments."""
    global _active_process
    _active_process = subprocess.Popen(
        [str(x) for x in args],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        start_new_session=True,
    )
    try:
        stdout, stderr = _active_process.communicate()
        if _active_process.returncode:
            message = stderr.decode(errors="replace")[-3500:]
            raise ConversionError(
                message or f"{Path(args[0]).name} failed ({_active_process.returncode})"
            )
        return stdout
    finally:
        _active_process = None


def unique(parent, stem, extension=""):
    candidate = parent / (stem + ("." + extension if extension else ""))
    i = 2
    while candidate.exists():
        candidate = parent / (f"{stem} {i}" + ("." + extension if extension else ""))
        i += 1
    return candidate


def image_open(path):
    if path.suffix.lower() == ".svg":
        doc = fitz.open(path)
        pix = doc[0].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=True)
        return Image.frombytes("RGBA", [pix.width, pix.height], pix.samples)
    with Image.open(path) as source:
        return ImageOps.exif_transpose(source).convert("RGBA")


def image_save(image, path, options):
    ext = path.suffix.lower().lstrip(".")
    fmt = {"jpg": "JPEG", "heic": "HEIF", "heif": "HEIF", "tif": "TIFF"}.get(ext, ext.upper())
    if ext in {"jpg", "jpeg", "bmp", "pdf"}:
        bg = Image.new("RGB", image.size, "white")
        if image.mode == "RGBA":
            bg.paste(image, mask=image.getchannel("A"))
        else:
            bg.paste(image)
        image = bg
    kwargs = (
        {"quality": max(1, min(100, int(options.get("quality", 80))))}
        if ext in {"jpg", "jpeg", "webp", "heic", "heif", "avif"}
        else {}
    )
    if ext == "png":
        kwargs.update(optimize=True)
        if options.get("metadata"):
            from PIL.PngImagePlugin import PngInfo

            tags = PngInfo()
            for key, value in options["metadata"].items():
                tags.add_text(key, str(value))
            kwargs["pnginfo"] = tags
    if options.get("metadata") and ext in {
        "jpg",
        "jpeg",
        "tiff",
        "webp",
        "heic",
        "avif",
    }:
        exif = Image.Exif()
        for key, value in options["metadata"].items():
            tag = {"title": 270, "author": 315, "copyright": 33432, "camera": 272}.get(key.lower())
            if tag:
                exif[tag] = value
        kwargs["exif"] = exif.tobytes()
    image.save(path, format=fmt, **kwargs)


def rectangle(options, image):
    # Region in normalized top-left coordinates; exact dimensions are optional.
    w, h = image.size
    x = float(options.get("x", 0))
    y = float(options.get("y", 0))
    rw = float(options.get("width", 1))
    rh = float(options.get("height", 1))
    if not (
        0 <= x < 1
        and 0 <= y < 1
        and 0 < rw <= 1
        and 0 < rh <= 1
        and x + rw <= 1.00001
        and y + rh <= 1.00001
    ):
        raise ConversionError("Crop region must be within the image (values from 0 to 1).")
    return (round(x * w), round(y * h), round((x + rw) * w), round((y + rh) * h))


def image_tool(srcs, output, tool, o):
    im = image_open(srcs[0])
    if tool == "removeBackground":
        prepared = o.get("preparedImages", {}).get(str(srcs[0]))
        if not prepared:
            raise ConversionError(
                "Background removal uses Apple Vision. Run it from the macOS app."
            )
        image_save(image_open(Path(prepared)), output, o)
        return
    if tool == "createPDF":
        doc = fitz.open()
        for src in srcs:
            pic = image_open(src)
            stream = io.BytesIO()
            pic.save(stream, format="PNG")
            page = doc.new_page(width=pic.width * 0.75, height=pic.height * 0.75)
            page.insert_image(page.rect, stream=stream.getvalue())
        doc.save(output)
        return
    if tool == "createCollage":
        pictures = [image_open(p) for p in srcs]
        layout = o.get("layout", "grid")
        count = len(pictures)
        columns = (
            count if layout == "row" else 1 if layout == "column" else math.ceil(math.sqrt(count))
        )
        rows = math.ceil(count / columns)
        cell = int(o.get("cellSize", 600))
        gap = int(o.get("spacing", 20))
        if cell < 1 or gap < 0 or columns * (cell + gap) > 30000 or rows * (cell + gap) > 30000:
            raise ConversionError("Collage dimensions are too large.")
        im = Image.new(
            "RGBA",
            (columns * cell + (columns + 1) * gap, rows * cell + (rows + 1) * gap),
            o.get("background", "#fff8ef"),
        )
        for i, pic in enumerate(pictures):
            tile = ImageOps.fit(pic, (cell, cell))
            if layout == "featured" and i == 0 and count > 2:
                # Featured uses a wide first image, followed by the remaining thumbnails.
                continue
            radius = int(o.get("radius", 0))
            if radius > 0:
                mask = Image.new("L", tile.size)
                ImageDraw.Draw(mask).rounded_rectangle((0, 0, cell, cell), radius=radius, fill=255)
                tile.putalpha(
                    Image.composite(tile.getchannel("A"), Image.new("L", tile.size), mask)
                )
            im.alpha_composite(
                tile,
                (
                    gap + (i % columns) * (cell + gap),
                    gap + (i // columns) * (cell + gap),
                ),
            )
        if layout == "featured" and count > 2:
            columns = count - 1
            im = Image.new(
                "RGBA",
                (1200 + gap * 2, 600 + int(1200 / columns) + gap * 3),
                o.get("background", "#fff8ef"),
            )
            im.alpha_composite(ImageOps.fit(pictures[0], (1200, 600)), (gap, gap))
            cell = (1200 - gap * (columns - 1)) // columns
            for i, pic in enumerate(pictures[1:]):
                im.alpha_composite(
                    ImageOps.fit(pic, (cell, cell)),
                    (gap + i * (cell + gap), 600 + 2 * gap),
                )
    elif tool in {"cropImage", "resizeImage"}:
        if tool == "cropImage":
            im = im.crop(rectangle(o, im))
        if tool == "resizeImage" and not (o.get("outputWidth") or o.get("outputHeight")):
            scale = float(o.get("scale", 1))
            if not 0 < scale <= 10:
                raise ConversionError("Scale must be between 0 and 10.")
            im = im.resize(
                (max(1, int(im.width * scale)), max(1, int(im.height * scale))),
                Image.Resampling.LANCZOS,
            )
        if o.get("outputWidth") or o.get("outputHeight"):
            w = int(
                o.get("outputWidth")
                or (
                    round(im.width * int(o["outputHeight"]) / im.height)
                    if o.get("outputHeight")
                    else im.width
                )
            )
            h = int(o.get("outputHeight") or round(im.height * w / im.width))
            if w < 1 or h < 1 or w * h > 150_000_000:
                raise ConversionError("Invalid output dimensions.")
            im = im.resize((w, h), Image.Resampling.LANCZOS)
    elif tool == "editImage":
        for name, enhancer, default in [
            ("brightness", ImageEnhance.Brightness, 1),
            ("contrast", ImageEnhance.Contrast, 1),
            ("saturation", ImageEnhance.Color, 1),
            ("sharpness", ImageEnhance.Sharpness, 1),
        ]:
            im = enhancer(im).enhance(float(o.get(name, default)))
        im = im.rotate(-float(o.get("rotation", 0)), expand=True, resample=Image.Resampling.BICUBIC)
        effect = o.get("effect", "original")
        alpha = im.getchannel("A")
        if effect == "grayscale":
            im = ImageOps.grayscale(im).convert("RGBA")
        elif effect == "sepia":
            im = ImageOps.colorize(ImageOps.grayscale(im), "#291c0e", "#ffe4b5").convert("RGBA")
        elif effect == "invert":
            im = ImageOps.invert(im.convert("RGB")).convert("RGBA")
        elif effect == "blur":
            im = im.filter(ImageFilter.GaussianBlur(float(o.get("blur", 5))))
        im.putalpha(alpha)
    elif tool == "annotateImage":
        draw = ImageDraw.Draw(im)
        color = o.get("color", "#ff702e")
        stroke = int(o.get("stroke", 6))
        try:
            font = ImageFont.truetype(
                "/System/Library/Fonts/Helvetica.ttc",
                int(o.get("fontSize", max(20, im.width // 20))),
            )
        except OSError:
            font = ImageFont.load_default()
        annotations = o.get("annotations") or [
            {
                "type": o.get("annotationType", "text"),
                "text": o.get("text", "Annotation"),
                "x": float(o.get("x", 0.1)),
                "y": float(o.get("y", 0.1)),
                "width": float(o.get("width", 0.6)),
                "height": float(o.get("height", 0.25)),
            }
        ]
        for a in annotations:
            x = float(a.get("x", 0.1)) * im.width
            y = float(a.get("y", 0.1)) * im.height
            w = float(a.get("width", 0.6)) * im.width
            h = float(a.get("height", 0.25)) * im.height
            kind = a.get("type", "text")
            if kind == "text":
                draw.text((x, y), a.get("text", "Annotation"), fill=color, font=font)
            elif kind == "rectangle":
                draw.rectangle((x, y, x + w, y + h), outline=color, width=stroke)
            elif kind == "ellipse":
                draw.ellipse((x, y, x + w, y + h), outline=color, width=stroke)
            elif kind == "callout":
                r = max(20, im.width // 30)
                draw.ellipse((x, y, x + 2 * r, y + 2 * r), fill=color)
                draw.text(
                    (x + r, y + r),
                    a.get("text", "1"),
                    fill="white",
                    font=font,
                    anchor="mm",
                )
            elif kind == "arrow":
                draw.line((x, y, x + w, y + h), fill=color, width=stroke)
                angle = math.atan2(h, w)
                size = max(18, stroke * 4)
                draw.polygon(
                    [
                        (x + w, y + h),
                        (
                            x + w - size * math.cos(angle - 0.5),
                            y + h - size * math.sin(angle - 0.5),
                        ),
                        (
                            x + w - size * math.cos(angle + 0.5),
                            y + h - size * math.sin(angle + 0.5),
                        ),
                    ],
                    fill=color,
                )
    elif tool == "frameImage":
        gap = int(o.get("spacing", 60))
        radius = int(o.get("radius", 24))
        w, h = im.size
        if gap < 0 or gap > 5000:
            raise ConversionError("Background spacing must be between 0 and 5000.")
        bg = Image.new("RGBA", (w + gap * 2, h + gap * 2), o.get("background", "#fff0df"))
        if o.get("backgroundImage"):
            bg = ImageOps.fit(image_open(Path(o["backgroundImage"])), bg.size)
        elif o.get("gradient", False):
            top = Image.new("RGBA", bg.size, o.get("color", "#ff702e"))
            mask = Image.linear_gradient("L").resize(bg.size)
            bg = Image.composite(bg, top, mask)
        mask = Image.new("L", im.size)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, w, h), radius=radius, fill=255)
        mask = Image.composite(im.getchannel("A"), Image.new("L", im.size), mask)
        bg.paste(im, (gap, gap), mask)
        im = bg
    elif tool == "redactImage":
        regions = o.get("regions") or [
            dict(
                x=o.get("x", 0.1),
                y=o.get("y", 0.1),
                width=o.get("width", 0.3),
                height=o.get("height", 0.2),
            )
        ]
        for region in regions:
            box = rectangle(region, im)
            if o.get("redaction", "solid") == "blur":
                im.paste(
                    im.crop(box).filter(ImageFilter.GaussianBlur(max(20, im.width / 40))),
                    box,
                )
            else:
                ImageDraw.Draw(im).rectangle(box, fill=o.get("color", "#000000"))
    elif tool == "compress":
        scale = float(o.get("scale", 1))
        if not 0 < scale <= 1:
            raise ConversionError("Scale must be between 0 and 1.")
        im.thumbnail(
            (max(1, int(im.width * scale)), max(1, int(im.height * scale))),
            Image.Resampling.LANCZOS,
        )
    elif tool not in {"convert", "removeMetadata"}:
        raise ConversionError("Unsupported image operation.")
    image_save(im, output, o)
    target = int(o.get("targetKB", 0)) * 1024
    if tool == "compress" and o.get("reduction"):
        target = int(srcs[0].stat().st_size * (1 - float(o["reduction"]) / 100))
    if tool == "compress" and target and output.stat().st_size > target:
        for quality in range(int(o.get("quality", 80)) - 10, 9, -10):
            image_save(im, output, {**o, "quality": quality})
            if output.stat().st_size <= target:
                break
        for _ in range(16):
            if output.stat().st_size <= target:
                break
            im = im.resize(
                (max(1, int(im.width * 0.85)), max(1, int(im.height * 0.85))),
                Image.Resampling.LANCZOS,
            )
            image_save(im, output, {**o, "quality": 20})
        if output.stat().st_size > target:
            raise ConversionError("The target size is too small for this format.")


def probe(path):
    return json.loads(
        run(
            [
                executable("ffprobe"),
                "-v",
                "error",
                "-show_format",
                "-show_streams",
                "-of",
                "json",
                path,
            ]
        )
    )


@lru_cache(maxsize=4)
def media_encoders(binary):
    listing = run([binary, "-hide_banner", "-encoders"]).decode(errors="replace")
    return frozenset(re.findall(r"^\s*[VAS][A-Z.]{5}\s+(\S+)", listing, re.MULTILINE))


def media_codecs(ext, audio_only=False):
    if audio_only:
        if ext == "ogg":
            available = media_encoders(executable("ffmpeg"))
            encoder = next((name for name in ("libvorbis", "libopus") if name in available), None)
            if encoder is None:
                raise ConversionError("This FFmpeg build needs a Vorbis or Opus encoder for OGG.")
            return ["-c:a", encoder]
        return {
            "mp3": ["-c:a", "libmp3lame"],
            "m4a": ["-c:a", "aac"],
            "wav": ["-c:a", "pcm_s16le"],
            "flac": ["-c:a", "flac"],
            "opus": ["-c:a", "libopus"],
            "aiff": ["-c:a", "pcm_s16be"],
            "wma": ["-c:a", "wmav2"],
        }[ext]
    return {
        "mp4": [
            "-c:v",
            "libx264",
            "-pix_fmt",
            "yuv420p",
            "-c:a",
            "aac",
            "-movflags",
            "+faststart",
        ],
        "mov": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
        "mkv": ["-c:v", "libx264", "-c:a", "aac"],
        "webm": ["-c:v", "libvpx-vp9", "-pix_fmt", "yuva420p", "-c:a", "libopus"],
        "avi": ["-c:v", "mpeg4", "-c:a", "libmp3lame"],
        "wmv": ["-c:v", "wmv2", "-c:a", "wmav2"],
        "gif": ["-an", "-loop", "0"],
    }[ext]


def media_tool(srcs, output, tool, o):
    ff = executable("ffmpeg")
    src = srcs[0]
    ext = output.suffix[1:].lower()
    audio_only = ext in AUDIO_EXTS
    args = [ff, "-hide_banner", "-loglevel", "error", "-nostdin", "-y", "-i", src]
    filters = []
    af = []
    info = probe(src)
    has_audio = any(s["codec_type"] == "audio" for s in info["streams"])
    if tool in {"trimAudio", "trimVideo"}:
        start = float(o.get("start", 0))
        end = float(o.get("end", 0))
        duration = float(info["format"].get("duration", 0))
        if start < 0 or start >= duration or (end and (end <= start or end > duration + 0.1)):
            raise ConversionError("Choose a valid start and end within the file duration.")
        args += ["-ss", str(start)]
        if end:
            args += ["-t", str(end - start)]
        if tool == "trimAudio" and o.get("removeSilence"):
            # Reverse and trim again: only the ends are removed, internal pauses are kept.
            af += [
                "silenceremove=start_periods=1:start_duration=0.05:start_threshold=-45dB",
                "areverse",
                "silenceremove=start_periods=1:start_duration=0.05:start_threshold=-45dB",
                "areverse",
            ]
    elif tool == "normalizeAudio":
        af += [f"loudnorm=I={float(o.get('lufs', -16))}:TP=-1.5:LRA=11"]
    elif tool == "audioChannels":
        args += ["-ac", str(int(o.get("channels", 1)))]
    elif tool == "muteVideo":
        args += ["-an"]
    elif tool == "changeVideoSpeed":
        speed = float(o.get("speed", 2))
        if not 0.25 <= speed <= 4:
            raise ConversionError("Speed must be between 0.25 and 4.")
        filters += [f"setpts=(PTS-STARTPTS)/{speed}"]
        remaining = speed
        while remaining > 2:
            af.append("atempo=2")
            remaining /= 2
        while remaining < 0.5:
            af.append("atempo=0.5")
            remaining /= 0.5
        af.append(f"atempo={remaining}")
    elif tool in {"cropVideo", "redactVideo"}:
        stream = next(s for s in info["streams"] if s["codec_type"] == "video")
        w, h = stream["width"], stream["height"]
        rotation = next(
            (
                item.get("rotation", 0)
                for item in stream.get("side_data_list", [])
                if "rotation" in item
            ),
            0,
        )
        if abs(int(rotation)) % 180 == 90:
            w, h = h, w
        box = rectangle(o, Image.new("L", (w, h)))
        x, y, x2, y2 = box
        cw = max(2, (x2 - x) // 2 * 2)
        ch = max(2, (y2 - y) // 2 * 2)
        if tool == "cropVideo":
            filters += [f"crop={cw}:{ch}:{x}:{y}"]
        else:
            enable = f":enable='between(t,{float(o.get('start', 0))},{float(o.get('end') or info['format']['duration'])})'"
            if o.get("redaction") == "blur":
                args += [
                    "-filter_complex",
                    rf"[0:v]split[base][region];[region]crop={cw}:{ch}:{x}:{y},boxblur=luma_radius=min(20\,min(w\,h)/4):luma_power=2:chroma_radius=min(10\,min(w\,h)/4):chroma_power=2[blur];[base][blur]overlay={x}:{y}"
                    + enable
                    + "[out]",
                    "-map",
                    "[out]",
                    "-map",
                    "0:a?",
                ]
            else:
                filters += [f"drawbox=x={x}:y={y}:w={cw}:h={ch}:color=black:t=fill" + enable]
    elif tool == "redactAudio":
        ranges = parse_times(o.get("ranges", "0:1"), ranges=True)
        duration = float(info["format"]["duration"])
        if any(a < 0 or b <= a or b > duration + 0.1 for a, b in ranges):
            raise ConversionError("Bleep ranges must be inside the recording.")
        enable = "+".join(f"between(t,{a},{b})" for a, b in ranges)
        args += [
            "-filter_complex",
            f"[0:a]volume=0:enable='{enable}'[clean];sine=frequency=1000:sample_rate=48000:duration={duration},volume=0.2,volume=0:enable='not({enable})'[tone];[clean][tone]amix=inputs=2:normalize=0[out]",
            "-map",
            "[out]",
        ]
    elif tool == "audioToVideo":
        width = int(o.get("outputWidth", 1280))
        height = int(o.get("outputHeight", 720))
        if width < 16 or height < 16 or width > 4096 or height > 4096:
            raise ConversionError("Visualizer dimensions must be between 16 and 4096.")
        if o.get("backgroundImage"):
            args += [
                "-loop",
                "1",
                "-i",
                o["backgroundImage"],
                "-map",
                "1:v",
                "-map",
                "0:a",
                "-vf",
                f"scale={width}:{height}:force_original_aspect_ratio=increase,crop={width}:{height}",
                "-shortest",
            ]
        else:
            args += [
                "-filter_complex",
                f"[0:a]showwaves=s={width}x{height}:mode=line:colors=0xff702e:rate=30,format=yuv420p[v]",
                "-map",
                "[v]",
                "-map",
                "0:a",
            ]
    elif tool in {"splitVideo", "videoSnapshots"}:
        output.mkdir()
        duration = float(info["format"]["duration"])
        times = parse_times(o.get("times", "1"))
        if any(t < 0 or t >= duration for t in times):
            raise ConversionError("Timestamps must be within the file duration.")
        if tool == "videoSnapshots":
            for i, t in enumerate(times):
                run(
                    [
                        ff,
                        "-hide_banner",
                        "-loglevel",
                        "error",
                        "-nostdin",
                        "-y",
                        "-ss",
                        str(t),
                        "-i",
                        src,
                        "-frames:v",
                        "1",
                        output / f"Frame {i + 1:03}.png",
                    ]
                )
        else:
            points = [0] + sorted(set(t for t in times if t > 0)) + [duration]
            for i, (a, b) in enumerate(zip(points, points[1:])):
                run(
                    [
                        ff,
                        "-hide_banner",
                        "-loglevel",
                        "error",
                        "-nostdin",
                        "-y",
                        "-i",
                        src,
                        "-ss",
                        str(a),
                        "-t",
                        str(b - a),
                        *media_codecs("mp4"),
                        output / f"Clip {i + 1:03}.mp4",
                    ]
                )
        return
    elif tool == "joinVideos":
        # Normalize resolution and audio on each input before concatenation.
        with tempfile.TemporaryDirectory(prefix="tangerine-join-") as temp:
            files = []
            for i, path in enumerate(srcs):
                p = Path(temp) / f"{i}.mp4"
                pinfo = probe(path)
                audio = any(s["codec_type"] == "audio" for s in pinfo["streams"])
                command = [
                    ff,
                    "-hide_banner",
                    "-loglevel",
                    "error",
                    "-nostdin",
                    "-y",
                    "-i",
                    path,
                ]
                if not audio:
                    command += [
                        "-f",
                        "lavfi",
                        "-i",
                        "anullsrc=r=48000:cl=stereo",
                        "-shortest",
                    ]
                command += [
                    "-map",
                    "0:v:0",
                    "-map",
                    "0:a:0" if audio else "1:a:0",
                    "-vf",
                    "scale=1280:720:force_original_aspect_ratio=decrease,pad=1280:720:(ow-iw)/2:(oh-ih)/2,setsar=1",
                    "-r",
                    "30",
                    "-ar",
                    "48000",
                    "-ac",
                    "2",
                    *media_codecs("mp4"),
                    p,
                ]
                run(command)
                files.append(p)
            listing = Path(temp) / "list.txt"
            listing.write_text("\n".join(f"file '{p}'" for p in files))
            run(
                [
                    ff,
                    "-hide_banner",
                    "-loglevel",
                    "error",
                    "-nostdin",
                    "-y",
                    "-f",
                    "concat",
                    "-safe",
                    "0",
                    "-i",
                    listing,
                    "-c",
                    "copy",
                    "-movflags",
                    "+faststart",
                    output,
                ]
            )
        return
    elif tool not in {"convert", "compress", "removeMetadata"}:
        raise ConversionError("Unsupported media operation.")
    if tool == "removeMetadata":
        args += ["-map_metadata", "-1", "-map_chapters", "-1"]
        for key, value in o.get("metadata", {}).items():
            args += ["-metadata", f"{key}={value}"]
    if audio_only:
        args += ["-vn"]
    if filters:
        args += ["-vf", ",".join(filters)]
    if af and has_audio:
        args += ["-af", ",".join(af)]
    args += media_codecs(ext, audio_only)
    if tool == "compress":
        target = float(o.get("targetKB", 0)) * 1024
        if o.get("reduction"):
            target = src.stat().st_size * (1 - float(o["reduction"]) / 100)
        if target > 0:
            duration = float(info["format"].get("duration", 0))
            if duration <= 0:
                raise ConversionError("Cannot calculate a target bitrate without duration.")
            budget = max(16, target * 8 / duration / 1000 * 0.95)
            if audio_only:
                o = {**o, "audioKbps": min(320, max(16, int(budget)))}
            else:
                audio_rate = (
                    min(int(o.get("audioKbps", 128)), max(32, int(budget * 0.2)))
                    if has_audio
                    else 0
                )
                o = {
                    **o,
                    "audioKbps": audio_rate,
                    "videoKbps": max(32, int(budget - audio_rate)),
                }
        if not audio_only:
            args += ["-b:v", f"{int(o.get('videoKbps', 1800))}k"]
        if ext not in {"wav", "flac", "aiff"}:
            args += ["-b:a", f"{int(o.get('audioKbps', 128))}k"]
    args += [output]
    run(args)


def parse_times(value, ranges=False):
    try:
        values = [v.strip() for v in str(value).split(",") if v.strip()]
        if not values:
            raise ValueError()
        if ranges:
            return [
                tuple(float(t) for t in v.split(":")) for v in values if len(v.split(":")) == 2
            ] or (_ for _ in ()).throw(ValueError())
        return [float(v) for v in values]
    except ValueError:
        raise ConversionError("Use comma-separated seconds, or start:end pairs for ranges.")


def pdf_tool(srcs, output, tool, o):
    doc = fitz.open(srcs[0])
    if doc.needs_pass:
        if not doc.authenticate(o.get("password", "")):
            raise ConversionError("This PDF is encrypted. Enter its password in Advanced options.")
    if tool == "splitPDF":
        output.mkdir()
        for i in range(len(doc)):
            part = fitz.open()
            part.insert_pdf(doc, from_page=i, to_page=i)
            part.save(output / f"Page {i + 1:03}.pdf")
    elif tool == "mergePDF":
        result = fitz.open()
        for p in srcs:
            other = fitz.open(p)
            if other.needs_pass and not other.authenticate(o.get("password", "")):
                raise ConversionError(f"Password required for {p.name}.")
            result.insert_pdf(other)
        result.save(output, garbage=4, deflate=True)
    elif tool == "organizePDF":
        order = str(o.get("pages") or ",".join(str(i + 1) for i in range(len(doc))))
        indexes = []
        try:
            for token in order.split(","):
                if "-" in token:
                    a, b = map(int, token.split("-"))
                    indexes.extend(range(a - 1, b))
                else:
                    indexes.append(int(token) - 1)
        except ValueError:
            raise ConversionError("Page order uses numbers and ranges, such as 3,1,2,2,4-6.")
        if not indexes or any(i < 0 or i >= len(doc) for i in indexes):
            raise ConversionError("Page number is outside the document.")
        result = fitz.open()
        selections = o.get("pageSelections")
        if selections is not None:
            if not selections:
                raise ConversionError("Keep at least one PDF page.")
            for item in selections:
                index = int(item["page"]) - 1
                rotation = int(item.get("rotation", 0))
                if index < 0 or index >= len(doc) or rotation % 90:
                    raise ConversionError("Invalid PDF page selection.")
                result.insert_pdf(doc, from_page=index, to_page=index)
                result[-1].set_rotation((result[-1].rotation + rotation) % 360)
        else:
            for index in indexes:
                result.insert_pdf(doc, from_page=index, to_page=index)
        rotation = int(o.get("rotation", 0))
        if rotation % 90:
            raise ConversionError("PDF rotation must be a multiple of 90 degrees.")
        for page in result:
            page.set_rotation((page.rotation + rotation) % 360)
        result.save(output, garbage=4, deflate=True)
    elif tool == "removeMetadata":
        doc.set_metadata(o.get("metadata", {}))
        doc.del_xml_metadata()
        doc.save(output, garbage=4, deflate=True)
    elif tool == "compress":
        # Re-encode embedded images, preserving selectable text and vectors.
        for page in doc:
            for item in page.get_images(full=True):
                xref = item[0]
                try:
                    data = doc.extract_image(xref)
                    im = Image.open(io.BytesIO(data["image"])).convert("RGB")
                    im.thumbnail((1600, 1600))
                    stream = io.BytesIO()
                    im.save(stream, format="JPEG", quality=int(o.get("quality", 65)))
                    if len(stream.getvalue()) < len(data["image"]):
                        page.replace_image(xref, stream=stream.getvalue())
                except (ValueError, RuntimeError, OSError):
                    pass
        doc.save(output, garbage=4, deflate=True)
    elif tool == "convert":
        ext = output.suffix[1:]
        if ext in {"png", "jpg"}:
            output.mkdir()
            for i, page in enumerate(doc):
                page.get_pixmap(dpi=300, alpha=False).save(output / f"Page {i + 1:03}.{ext}")
        else:
            pdf_export(doc, output, ext, o)
    else:
        raise ConversionError("Unsupported PDF operation.")


def pdf_export(doc, output, ext, o):
    texts = o.get("textOverride") or [p.get_text(sort=True) for p in doc]
    text = "\n\n".join(texts)
    if ext in {"txt", "md"}:
        output.write_text(text, encoding="utf-8")
    elif ext == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        deck.slide_width = Inches(12)
        deck.slide_height = int(deck.slide_width * doc[0].rect.height / doc[0].rect.width)
        for page in doc:
            pix = page.get_pixmap(dpi=150)
            stream = io.BytesIO(pix.tobytes("png"))
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(stream, 0, 0, width=deck.slide_width, height=deck.slide_height)
        deck.save(output)
    elif ext == "epub" and o.get("epubLayout") == "fixed":
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier("tangerine-fixed-" + output.stem)
        book.set_title(output.stem)
        book.set_language("en")
        book.add_metadata(None, "meta", "pre-paginated", {"property": "rendition:layout"})
        chapters = []
        for index, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150, alpha=False)
            image = epub.EpubItem(
                uid=f"image-{index}",
                file_name=f"images/page-{index + 1}.png",
                media_type="image/png",
                content=pix.tobytes("png"),
            )
            book.add_item(image)
            chapter = epub.EpubHtml(title=f"Page {index + 1}", file_name=f"page-{index + 1}.xhtml")
            chapter.content = f'<html xmlns="http://www.w3.org/1999/xhtml"><head><meta name="viewport" content="width={pix.width},height={pix.height}"/></head><body style="margin:0"><img src="images/page-{index + 1}.png" style="width:100%;height:100%" alt="Page {index + 1}"/></body></html>'
            book.add_item(chapter)
            chapters.append(chapter)
        book.toc = chapters
        book.spine = ["nav"] + chapters
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(str(output), book)
    elif ext == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        for i, page in enumerate(doc):
            sheet = wb.create_sheet(f"Page {i + 1}")
            tables = page.find_tables().tables
            if tables:
                for table in tables:
                    for row in table.extract():
                        sheet.append([safe_cell(v) for v in row])
            else:
                for line in texts[i].splitlines():
                    sheet.append([safe_cell(v) for v in re.split(r"\t| {2,}", line)])
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 28
        wb.save(output)
    else:
        write_document(text, output, ext, o)


def safe_cell(value):
    # Imported text is data, never an executable spreadsheet formula.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def office(src, output, ext):
    soffice = executable("soffice")
    with tempfile.TemporaryDirectory(prefix="tangerine-office-") as temp:
        folder = Path(temp)
        filters = {
            "csv": "csv:Text - txt - csv (StarCalc):44,34,76",
            "tsv": "csv:Text - txt - csv (StarCalc):9,34,76",
        }
        command = [
            soffice,
            f"-env:UserInstallation={(folder / 'profile').as_uri()}",
            "--headless",
            "--convert-to",
            filters.get(ext, ext),
            "--outdir",
            folder,
            src,
        ]
        run(command)
        expected = list(folder.glob("*." + ("csv" if ext == "tsv" else ext)))
        if not expected:
            raise ConversionError(f"LibreOffice could not convert {src.name} to {ext.upper()}.")
        shutil.copyfile(expected[0], output)


def pages_application_id():
    # Both editions ship Apple's Pages AppleScript dictionary under different IDs.
    for root in (Path("/Applications"), Path.home() / "Applications"):
        for name in ("Pages.app", "Pages Creator Studio.app"):
            info = root / name / "Contents/Info.plist"
            try:
                with info.open("rb") as handle:
                    bundle_id = plistlib.load(handle).get("CFBundleIdentifier")
                if bundle_id in {"com.apple.iWork.Pages", "com.apple.Pages"}:
                    return bundle_id
            except (OSError, ValueError, plistlib.InvalidFileException):
                continue
    return "com.apple.iWork.Pages"


def pages_export(src, output):
    # Apple Pages owns the .pages format. Apple Events permission is requested by macOS.
    script = """on run argv
set inputFile to POSIX file (item 1 of argv)
set outputFile to POSIX file (item 2 of argv)
tell application id "com.apple.iWork.Pages"
set d to open inputFile
save d in outputFile
close d saving no
end tell
end run"""
    run(
        [
            "/usr/bin/osascript",
            "-e",
            script.replace("com.apple.iWork.Pages", pages_application_id()),
            src,
            output,
        ]
    )


def pages_import(src, output):
    # LaunchServices sends a document-open event with a sandbox extension. An
    # AppleScript file spec alone can leave sandboxed Pages unable to read a
    # selected document outside its container ("Operation not permitted").
    run(["/usr/bin/open", "-b", pages_application_id(), src])
    script = """on run argv
tell application id "com.apple.iWork.Pages"
set d to open (POSIX file (item 1 of argv))
export d to (POSIX file (item 2 of argv)) as Microsoft Word
close d saving no
end tell
end run"""
    run(
        [
            "/usr/bin/osascript",
            "-e",
            script.replace("com.apple.iWork.Pages", pages_application_id()),
            src,
            output,
        ]
    )


def read_text_clipping(src):
    import struct

    try:
        resource = Path(str(src) + "/..namedfork/rsrc").read_bytes()
    except OSError:
        raise ConversionError(
            "This text clipping has no readable macOS resource fork. Copy it in Finder with metadata intact."
        )
    try:
        data_offset, map_offset, _, _ = struct.unpack_from(">4I", resource, 0)
        type_offset = struct.unpack_from(">H", resource, map_offset + 24)[0]
        types = map_offset + type_offset
        count = struct.unpack_from(">H", resource, types)[0] + 1
        found = {}
        for i in range(count):
            kind, number, reference = struct.unpack_from(">4sHH", resource, types + 2 + i * 8)
            if kind not in {b"utxt", b"utf8", b"TEXT", b"RTF "}:
                continue
            offset = types + reference
            record = resource[offset : offset + 12]
            position = data_offset + int.from_bytes(record[5:8], "big")
            size = struct.unpack_from(">I", resource, position)[0]
            if size > len(resource) - position - 4:
                raise ValueError()
            found[kind] = resource[position + 4 : position + 4 + size]
        if b"utf8" in found:
            return found[b"utf8"].decode("utf-8")
        if b"utxt" in found:
            return found[b"utxt"].decode("utf-16-be")
        if b"TEXT" in found:
            return found[b"TEXT"].decode("mac_roman")
        if b"RTF " in found:
            with tempfile.TemporaryDirectory() as temp:
                rtf = Path(temp) / "clip.rtf"
                rtf.write_bytes(found[b"RTF "])
                return run(["/usr/bin/textutil", "-convert", "txt", "-stdout", rtf]).decode("utf-8")
        raise ValueError()
    except (ValueError, struct.error, UnicodeError, IndexError):
        raise ConversionError("The text clipping does not contain valid text resources.")


def read_document(src):
    ext = src.suffix.lower()[1:]
    if ext == "docx":
        from docx import Document

        d = Document(src)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            parts.extend("\t".join(c.text for c in row.cells) for row in table.rows)
        return "\n".join(parts)
    if ext == "epub":
        from bs4 import BeautifulSoup
        from ebooklib import ITEM_DOCUMENT, epub

        book = epub.read_epub(str(src), options={"ignore_ncx": True})
        items = {item.id: item for item in book.get_items_of_type(ITEM_DOCUMENT)}
        return "\n\n".join(
            BeautifulSoup(items[identifier].get_content(), "html.parser").get_text("\n")
            for identifier, _ in book.spine
            if identifier in items
        )
    if ext == "html":
        from bs4 import BeautifulSoup

        return BeautifulSoup(src.read_text(encoding="utf-8"), "html.parser").get_text("\n")
    if ext == "textclipping":
        return read_text_clipping(src)
    if ext in {"rtf", "rtfd"}:
        return run(["/usr/bin/textutil", "-convert", "txt", "-stdout", src]).decode("utf-8")
    return src.read_text(encoding="utf-8-sig")


def write_document(text, output, ext, o):
    if ext in {"txt", "md"}:
        output.write_text(text, encoding="utf-8")
    elif ext == "html":
        import markdown

        output.write_text(
            '<!doctype html><html><head><meta charset="utf-8"><title>'
            + html.escape(output.stem)
            + "</title></head><body>"
            + markdown.markdown(text, extensions=["tables", "fenced_code"])
            + "</body></html>",
            encoding="utf-8",
        )
    elif ext == "docx":
        from docx import Document

        doc = Document()
        for line in text.splitlines():
            match = re.match(r"^(#{1,6})\s+(.*)", line)
            if match:
                doc.add_heading(match[2], len(match[1]))
            elif "\t" in line:
                values = line.split("\t")
                cells = doc.add_table(rows=1, cols=len(values)).rows[0].cells
                for cell, value in zip(cells, values):
                    cell.text = value
            else:
                doc.add_paragraph(line)
        doc.save(output)
    elif ext == "epub":
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier("tangerine-" + output.stem)
        book.set_title(output.stem)
        book.set_language("en")
        chapter = epub.EpubHtml(title=output.stem, file_name="content.xhtml", lang="en")
        chapter.content = (
            "<html><body>"
            + "".join("<p>" + html.escape(line) + "</p>" for line in text.splitlines())
            + "</body></html>"
        )
        book.add_item(chapter)
        book.toc = (chapter,)
        book.spine = ["nav", chapter]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(str(output), book)
    elif ext == "pages":
        with tempfile.TemporaryDirectory() as temp:
            docx = Path(temp) / "document.docx"
            write_document(text, docx, "docx", o)
            pages_export(docx, output)
    elif ext == "pdf":
        doc = fitz.open()
        story = fitz.Story(
            html="<html><body>"
            + "".join("<p>" + html.escape(line) + "</p>" for line in text.splitlines())
            + "</body></html>"
        )
        writer = fitz.DocumentWriter(str(output))
        more = True
        while more:
            device = writer.begin_page(fitz.Rect(0, 0, 595, 842))
            more, _ = story.place(fitz.Rect(42, 42, 553, 800))
            story.draw(device)
            writer.end_page()
        writer.close()
    else:
        raise ConversionError(f"Document output {ext} is unsupported.")


def document_convert(src, output, o):
    ext = output.suffix[1:]
    source = src.suffix.lower()[1:]
    with tempfile.TemporaryDirectory(prefix="tangerine-doc-") as temp:
        if source == "pages":
            converted = Path(temp) / "source.docx"
            pages_import(src, converted)
            src = converted
            source = "docx"
        if ext == "pages":
            if source == "docx":
                pages_export(src, output)
            else:
                write_document(read_document(src), output, ext, o)
        elif ext in {"png", "jpg"}:
            pdf = Path(temp) / "pages.pdf"
            document_convert(src, pdf, o)
            pdf_tool([pdf], output, "convert", o)
        elif source in {"ppt", "pptx", "xlsx", "xls", "ods", "csv", "tsv"} or (
            source == "docx" and ext in {"pdf", "html"}
        ):
            office(src, output, ext)
        else:
            write_document(read_document(src), output, ext, o)


def subtitle_convert(src, output):
    text = src.read_text(encoding="utf-8-sig").replace("\r\n", "\n")
    ext = output.suffix[1:]
    if ext == "txt":
        blocks = re.split(r"\n\s*\n", text)
        lines = []
        for block in blocks:
            part = block.splitlines()
            marker = next((i for i, line in enumerate(part) if "-->" in line), None)
            if marker is not None:
                lines.extend(re.sub("<[^>]+>", "", line) for line in part[marker + 1 :])
        text = "\n".join(lines) + "\n"
    elif ext == "vtt":
        text = "WEBVTT\n\n" + re.sub(r"(?m)^\d+\s*\n(?=\d\d:)", "", text)
        text = re.sub(r"(\d{2}:\d{2}:\d{2}),(\d{3})", r"\1.\2", text)
    elif ext == "srt":
        cues = []
        for block in re.split(r"\n\s*\n", text):
            lines = block.splitlines()
            i = next((i for i, line in enumerate(lines) if "-->" in line), None)
            if i is None:
                continue
            times = re.match(
                r"((?:\d{2}:)?\d{2}:\d{2}\.\d{3})\s*-->\s*((?:\d{2}:)?\d{2}:\d{2}\.\d{3})",
                lines[i],
            )
            if not times:
                continue
            a, b = [v if v.count(":") == 2 else "00:" + v for v in times.groups()]
            cues.append(
                f"{len(cues) + 1}\n{a.replace('.', ',')} --> {b.replace('.', ',')}\n"
                + "\n".join(lines[i + 1 :])
            )
        text = "\n\n".join(cues) + "\n"
    output.write_text(text, encoding="utf-8")


def safe_member(root, name):
    path = Path(name)
    if path.is_absolute() or ".." in path.parts or "\\" in name:
        raise ConversionError("Archive contains an unsafe path.")
    target = (root / path).resolve()
    if not target.is_relative_to(root.resolve()):
        raise ConversionError("Archive path escapes the destination.")
    return target


def extract(src, folder):
    limit = 20 * 1024**3
    count = 0
    total = 0
    if zipfile.is_zipfile(src):
        with zipfile.ZipFile(src) as archive:
            for member in archive.infolist():
                safe_member(folder, member.filename)
                if ((member.external_attr >> 16) & 0o170000) == 0o120000:
                    raise ConversionError("Symbolic links in archives are not extracted.")
                count += 1
                total += member.file_size
                if total > limit or count > 100000:
                    raise ConversionError(
                        "Archive exceeds the safety limit (20 GB / 100,000 entries)."
                    )
            archive.extractall(folder)
    elif tarfile.is_tarfile(src):
        with tarfile.open(src) as archive:
            for member in archive:
                safe_member(folder, member.name)
                if not (member.isdir() or member.isfile()):
                    raise ConversionError("Archive contains a link or special file.")
                count += 1
                total += member.size
                if total > limit or count > 100000:
                    raise ConversionError("Archive exceeds extraction limits.")
                archive.extract(member, folder, filter="data")
    elif src.suffix.lower() == ".gz":
        import gzip

        with gzip.open(src, "rb") as inp, open(folder / src.stem, "wb") as out:
            while chunk := inp.read(1024 * 1024):
                total += len(chunk)
                if total > limit:
                    raise ConversionError("Archive exceeds extraction limits.")
                out.write(chunk)
    elif src.suffix.lower() == ".rar":
        extract_rar(src, folder, limit)
    else:
        raise ConversionError("This is not a supported archive.")


def extract_rar(src, folder, limit):
    """Stream libarchive entries without permitting links or unbounded extraction."""
    library = (
        "/usr/lib/libarchive.2.dylib"
        if sys.platform == "darwin"
        else ctypes.util.find_library("archive")
    )
    if not library:
        raise ConversionError("System libarchive is unavailable.")
    lib = ctypes.CDLL(library)
    signatures = {
        "archive_read_new": ([], ctypes.c_void_p),
        "archive_read_support_filter_all": ([ctypes.c_void_p], ctypes.c_int),
        "archive_read_support_format_all": ([ctypes.c_void_p], ctypes.c_int),
        "archive_read_open_filename": (
            [ctypes.c_void_p, ctypes.c_char_p, ctypes.c_size_t],
            ctypes.c_int,
        ),
        "archive_read_next_header": (
            [ctypes.c_void_p, ctypes.POINTER(ctypes.c_void_p)],
            ctypes.c_int,
        ),
        "archive_entry_pathname": ([ctypes.c_void_p], ctypes.c_char_p),
        "archive_entry_filetype": ([ctypes.c_void_p], ctypes.c_uint),
        "archive_entry_hardlink": ([ctypes.c_void_p], ctypes.c_char_p),
        "archive_entry_size": ([ctypes.c_void_p], ctypes.c_int64),
        "archive_read_data": (
            [ctypes.c_void_p, ctypes.c_void_p, ctypes.c_size_t],
            ctypes.c_ssize_t,
        ),
        "archive_read_free": ([ctypes.c_void_p], ctypes.c_int),
        "archive_error_string": ([ctypes.c_void_p], ctypes.c_char_p),
    }
    for name, (args, result) in signatures.items():
        fn = getattr(lib, name)
        fn.argtypes = args
        fn.restype = result
    archive = lib.archive_read_new()
    entry = ctypes.c_void_p()
    buffer = ctypes.create_string_buffer(1024 * 1024)
    total = 0
    count = 0

    def error():
        return ConversionError(
            (lib.archive_error_string(archive) or b"Unable to read archive").decode(
                errors="replace"
            )
        )

    try:
        lib.archive_read_support_filter_all(archive)
        lib.archive_read_support_format_all(archive)
        if lib.archive_read_open_filename(archive, os.fsencode(src), 10240) != 0:
            raise error()
        while True:
            status = lib.archive_read_next_header(archive, ctypes.byref(entry))
            if status == 1:
                break
            if status != 0:
                raise error()
            name = lib.archive_entry_pathname(entry)
            if not name:
                raise ConversionError("Archive entry has no name.")
            path = safe_member(folder, os.fsdecode(name))
            kind = lib.archive_entry_filetype(entry)
            if kind not in {0o100000, 0o040000} or lib.archive_entry_hardlink(entry):
                raise ConversionError("Archive contains a link or special file.")
            count += 1
            if count > 100000 or lib.archive_entry_size(entry) > limit - total:
                raise ConversionError("Archive exceeds extraction limits.")
            if kind == 0o040000:
                path.mkdir(parents=True, exist_ok=True)
                continue
            path.parent.mkdir(parents=True, exist_ok=True)
            with path.open("xb") as target:
                while True:
                    size = lib.archive_read_data(archive, buffer, len(buffer))
                    if size < 0:
                        raise error()
                    if size == 0:
                        break
                    total += size
                    if total > limit:
                        raise ConversionError("Archive exceeds extraction limits.")
                    target.write(buffer.raw[:size])
    finally:
        lib.archive_read_free(archive)


def publish(staged, parent, stem, suffix):
    """Atomically publish without replacing a racing output from another process."""
    while True:
        out = unique(parent, stem, suffix)
        if sys.platform == "darwin":
            libc = ctypes.CDLL(None, use_errno=True)
            rename = libc.renamex_np
            rename.argtypes = [ctypes.c_char_p, ctypes.c_char_p, ctypes.c_uint]
            rename.restype = ctypes.c_int
            if rename(os.fsencode(staged), os.fsencode(out), 4) == 0:
                return out
            code = ctypes.get_errno()
            if code == errno.EEXIST:
                continue
            raise OSError(code, os.strerror(code))
        elif staged.is_file():
            try:
                os.link(staged, out)
                staged.unlink()
                return out
            except FileExistsError:
                continue
        else:
            # The application targets macOS; this fallback is for portable fixture tests.
            if out.exists():
                continue
            os.rename(staged, out)
            return out


def archive_tool(src, output, tool):
    if tool == "extractArchive":
        output.mkdir()
        extract(src, output)
        return
    with tempfile.TemporaryDirectory(prefix="tangerine-archive-") as temp:
        root = Path(temp)
        extract(src, root)
        ext = output.suffix[1:]
        if ext == "zip":
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                for path in sorted(root.rglob("*")):
                    archive.write(path, path.relative_to(root))
        else:
            with tarfile.open(output, "w:gz" if ext == "gz" else "w") as archive:
                for path in root.iterdir():
                    archive.add(path, arcname=path.name, recursive=True)


def execute(request):
    srcs = [Path(p).resolve() for p in request["inputs"]]
    if not srcs or any(not p.exists() for p in srcs):
        raise ConversionError("Select existing input files.")
    base_options = dict(request.get("options", {}))
    # Foundation can shorten /private/tmp to /tmp while pathlib resolves the other
    # way. Canonicalize both sides of cross-process lookup maps in one runtime.
    for key in ("preparedImages", "textOverrides"):
        if key in base_options:
            base_options[key] = {
                str(Path(path).resolve()): value for path, value in base_options[key].items()
            }
    tool = request.get("operation", "convert")
    fmt = request.get("format", "")
    parent = Path(request.get("destination") or srcs[0].parent).resolve()
    if not parent.is_dir():
        raise ConversionError("The output folder does not exist.")
    combined = tool in {"createPDF", "mergePDF", "joinVideos", "createCollage"}
    groups = [srcs] if combined else [[p] for p in srcs]
    outputs = []
    errors = []
    for group in groups:
        try:
            o = dict(base_options)
            if str(group[0]) in o.get("textOverrides", {}):
                o["textOverride"] = o["textOverrides"][str(group[0])]
            src = group[0]
            source = src.suffix.lower()[1:]
            ext = (
                fmt
                if tool == "convert"
                else (
                    "pdf"
                    if tool in {"createPDF", "mergePDF", "organizePDF", "splitPDF"}
                    else "mp4"
                    if tool
                    in {
                        "joinVideos",
                        "audioToVideo",
                        "cropVideo",
                        "changeVideoSpeed",
                        "redactVideo",
                        "trimVideo",
                        "muteVideo",
                    }
                    else "png"
                    if tool in {"annotateImage", "createCollage", "removeBackground"}
                    else source
                )
            )
            if tool == "compress" and source in {"wav", "aiff", "aif"}:
                ext = "flac"
            if tool in {"compress", "removeMetadata"} and source == "svg":
                ext = "png"
            if ext in {"jpeg", "tif", "heif", "aif"}:
                ext = {"jpeg": "jpg", "tif": "tiff", "heif": "heic", "aif": "aiff"}[ext]
            is_folder = tool in {
                "splitPDF",
                "extractArchive",
                "splitVideo",
                "videoSnapshots",
            } or (tool == "convert" and ext in {"jpg", "png"} and source not in IMAGE_EXTS)
            suffix = "" if is_folder else "tar.gz" if ext == "gz" else ext
            if not re.fullmatch("[a-z0-9]*", ext):
                raise ConversionError("Invalid output format.")
            stem = src.stem + ("" if tool == "convert" else " " + tool)
            out = unique(parent, stem, suffix)
            # All writes are staged on the destination volume, then atomically renamed.
            with tempfile.TemporaryDirectory(prefix=".tangerine-", dir=parent) as temp:
                staged = Path(temp) / ("output" + ("." + ext if ext else ""))
                if source in IMAGE_EXTS:
                    image_tool(group, staged, tool, o)
                elif source in AUDIO_EXTS | VIDEO_EXTS:
                    media_tool(group, staged, tool, o)
                elif source == "pdf":
                    pdf_tool(group, staged, tool, o)
                elif source in {"srt", "vtt"}:
                    subtitle_convert(src, staged)
                elif source in ARCHIVE_EXTS:
                    archive_tool(src, staged, tool)
                else:
                    if tool != "convert":
                        raise ConversionError("This operation does not support the selected file.")
                    document_convert(src, staged, o)
                if not staged.exists():
                    raise ConversionError("The converter produced no output.")
                if staged.is_file() and staged.stat().st_size == 0:
                    raise ConversionError("The converter produced an empty file.")
                # Recheck after long conversions; never replace another job's output.
                out = publish(staged, parent, stem, suffix)
            outputs.append(str(out))
        except Exception as error:
            if len(groups) == 1:
                raise
            errors.append({"input": str(group[0]), "error": str(error)})
    return {"outputs": outputs, **({"errors": errors} if errors else {})}


def main():
    try:
        request = json.load(open(sys.argv[1])) if len(sys.argv) > 1 else json.load(sys.stdin)
        from contextlib import redirect_stdout

        with redirect_stdout(sys.stderr):
            result = execute(request)
        print(json.dumps(result))
    except Exception as error:
        print(json.dumps({"error": str(error)}))
        return 1
    return 0


if __name__ == "__main__":
    signal.signal(signal.SIGTERM, cancel_worker)
    signal.signal(signal.SIGINT, cancel_worker)
    sys.exit(main())

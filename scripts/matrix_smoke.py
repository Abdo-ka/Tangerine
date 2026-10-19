#!/usr/bin/env python3
"""Exercise the reference's 275 format routes with generated, non-sensitive fixtures.

Apple Pages' 16 routes require interactive Automation permission and are reported
as manual checks. All other routes run locally. Report is written under .build/.
"""

import importlib.util
import json
import struct
import sys
import tempfile
import zipfile
import zlib
from pathlib import Path

import pymupdf as fitz
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("worker", ROOT / "Resources/engine.py")
e = importlib.util.module_from_spec(spec)
spec.loader.exec_module(e)
IMAGES = ["jpg", "png", "webp", "heic", "tiff", "svg", "avif", "bmp"]
AUDIO = ["mp3", "m4a", "wav", "flac", "ogg", "opus", "aiff", "wma"]
VIDEO = ["mp4", "mov", "mkv", "webm", "avi", "wmv"]
DOCUMENTS = ["docx", "pages", "md", "rtf", "rtfd", "textclipping", "txt"]
SHEETS = ["xlsx", "xls", "ods", "csv", "tsv"]
ARCHIVES = ["zip", "tar", "gz", "rar"]
INPUTS = (
    IMAGES
    + AUDIO
    + VIDEO
    + ["gif", "pdf"]
    + DOCUMENTS
    + ["ppt", "pptx"]
    + SHEETS
    + ["epub", "srt", "vtt"]
    + ARCHIVES
)


def formats(source):
    if source in IMAGES:
        outputs = ["jpg", "png", "webp", "heic", "tiff", "avif", "bmp", "pdf"]
    elif source in AUDIO:
        outputs = AUDIO
    elif source in VIDEO:
        outputs = VIDEO + ["gif", "mp3"]
    elif source == "gif":
        outputs = VIDEO
    elif source == "pdf":
        outputs = ["docx", "pages", "pptx", "xlsx", "jpg", "png", "epub", "txt", "md"]
    elif source in DOCUMENTS:
        outputs = ["docx", "pages", "pdf", "txt", "md", "html", "epub", "png", "jpg"]
    elif source in {"ppt", "pptx"}:
        outputs = ["pdf", "jpg", "png"]
    elif source in SHEETS:
        outputs = ["xlsx", "ods", "csv", "tsv", "pdf"]
    elif source == "epub":
        outputs = ["pdf", "docx", "pages", "txt", "md", "html"]
    elif source in {"srt", "vtt"}:
        outputs = ["srt", "vtt", "txt"]
    else:
        outputs = ["zip", "tar", "gz"]
    return [target for target in outputs if target != source]


def rar_fixture(path):
    """Generate a minimal RAR4 archive with a single stored text file."""

    def header(kind, flags, body):
        value = struct.pack("<BHH", kind, flags, 7 + len(body)) + body
        return struct.pack("<H", zlib.crc32(value) & 0xFFFF) + value

    name = b"fixture.txt"
    data = b"Local archive fixture\n"
    file_body = (
        struct.pack(
            "<IIBIIBBHI",
            len(data),
            len(data),
            3,
            zlib.crc32(data),
            0,
            20,
            0x30,
            len(name),
            0o100644,
        )
        + name
    )
    path.write_bytes(
        b"Rar!\x1a\x07\x00"
        + header(0x73, 0, b"\x00" * 6)
        + header(0x74, 0x8000, file_body)
        + data
        + header(0x7B, 0, b"")
    )


def validate(path, target):
    if path.is_dir():
        children = list(path.iterdir())
        assert children, "Empty output folder"
        for child in children:
            if child.suffix.lower() in {".png", ".jpg"}:
                with Image.open(child) as image:
                    image.verify()
        return
    assert path.stat().st_size > 0, "Empty output file"
    if target in {"jpg", "png", "webp", "heic", "tiff", "avif", "bmp"}:
        with Image.open(path) as image:
            image.verify()
    elif target == "pdf":
        with fitz.open(path) as document:
            assert len(document) > 0
    elif target in AUDIO + VIDEO + ["gif"]:
        assert e.probe(path)["streams"]
    elif target in {"docx", "pptx", "xlsx", "epub", "zip"}:
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
    elif target in {"tar", "gz"}:
        with e.tarfile.open(path) as archive:
            assert archive.getmembers()
    elif target in {"txt", "md", "html", "csv", "tsv", "srt", "vtt"}:
        assert path.read_text(encoding="utf-8-sig").strip()


def main():
    total = sum(len(formats(source)) for source in INPUTS)
    assert total == 275, total
    results = []
    with tempfile.TemporaryDirectory(prefix="tangerine-matrix-") as temporary:
        root = Path(temporary)
        source_dir = root / "inputs"
        source_dir.mkdir()
        dest = root / "outputs"
        dest.mkdir()
        sources = {}

        def convert(path, target):
            response = e.execute(
                {"inputs": [str(path)], "format": target, "destination": str(dest)}
            )
            return Path(response["outputs"][0])

        base = source_dir / "photo.png"
        Image.new("RGB", (96, 64), "#e86d24").save(base)
        for fmt in IMAGES:
            if fmt == "svg":
                path = source_dir / "vector.svg"
                path.write_text(
                    '<svg xmlns="http://www.w3.org/2000/svg" width="96" height="64"><rect width="96" height="64" fill="orange"/></svg>'
                )
            else:
                path = base if fmt == "png" else convert(base, fmt)
            sources[fmt] = path
        wav = source_dir / "tone.wav"
        video = source_dir / "clip.mp4"
        ff = e.executable("ffmpeg")
        e.run(
            [
                ff,
                "-hide_banner",
                "-loglevel",
                "error",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:sample_rate=48000",
                "-t",
                "0.5",
                wav,
            ]
        )
        e.run(
            [
                ff,
                "-hide_banner",
                "-loglevel",
                "error",
                "-f",
                "lavfi",
                "-i",
                "testsrc2=size=96x64:rate=10",
                "-f",
                "lavfi",
                "-i",
                "sine=sample_rate=48000",
                "-t",
                "0.5",
                "-c:v",
                "libx264",
                "-pix_fmt",
                "yuv420p",
                "-c:a",
                "aac",
                video,
            ]
        )
        for fmt in AUDIO:
            sources[fmt] = wav if fmt == "wav" else convert(wav, fmt)
        for fmt in VIDEO:
            sources[fmt] = video if fmt == "mp4" else convert(video, fmt)
        sources["gif"] = convert(video, "gif")
        pdf = source_dir / "pages.pdf"
        document = fitz.open()
        page = document.new_page()
        page.insert_text((40, 80), "Local conversion fixture")
        document.save(pdf)
        sources["pdf"] = pdf
        md = source_dir / "document.md"
        md.write_text("# Local fixture\n\nText for offline conversions.\n")
        sources["md"] = md
        sources["txt"] = convert(md, "txt")
        sources["docx"] = convert(md, "docx")
        sources["epub"] = convert(md, "epub")
        rtf = source_dir / "document.rtf"
        rtf.write_text(r"{\rtf1\ansi Local rich text fixture.}")
        sources["rtf"] = rtf
        rtfd = source_dir / "document.rtfd"
        e.run(["/usr/bin/textutil", "-convert", "rtfd", "-output", rtfd, rtf])
        sources["rtfd"] = rtfd
        # A real resource-fork clipping, rather than an incorrectly renamed text file.
        data = "Local Finder clipping".encode("utf-16-be")
        offset = 260 + len(data)
        resource = bytearray(offset + 50)
        struct.pack_into(">4I", resource, 0, 256, offset, 4 + len(data), 50)
        struct.pack_into(">I", resource, 256, len(data))
        resource[260:offset] = data
        struct.pack_into(">H", resource, offset + 24, 28)
        struct.pack_into(">H", resource, offset + 28, 0)
        struct.pack_into(">4sHH", resource, offset + 30, b"utxt", 0, 10)
        clipping = source_dir / "note.textClipping"
        clipping.write_bytes(b"")
        Path(str(clipping) + "/..namedfork/rsrc").write_bytes(resource)
        sources["textclipping"] = clipping
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Local slide"
        pptx = source_dir / "slides.pptx"
        presentation.save(pptx)
        sources["pptx"] = pptx
        ppt = source_dir / "slides.ppt"
        e.office(pptx, ppt, "ppt")
        sources["ppt"] = ppt
        csv = source_dir / "sheet.csv"
        csv.write_text("Name,Count\nOrange,3\n")
        sources["csv"] = csv
        for fmt in ["xlsx", "ods", "tsv"]:
            sources[fmt] = convert(csv, fmt)
        xls = source_dir / "sheet.xls"
        e.office(sources["xlsx"], xls, "xls")
        sources["xls"] = xls
        srt = source_dir / "captions.srt"
        srt.write_text("1\n00:00:00,100 --> 00:00:00,400\nLocal caption\n")
        sources["srt"] = srt
        sources["vtt"] = convert(srt, "vtt")
        archive = source_dir / "files.zip"
        with zipfile.ZipFile(archive, "w") as output:
            output.writestr("fixture.txt", "Local archive fixture")
        sources["zip"] = archive
        for fmt in ["tar", "gz"]:
            sources[fmt] = convert(archive, fmt)
        rar = source_dir / "files.rar"
        rar_fixture(rar)
        sources["rar"] = rar
        for source in INPUTS:
            for target in formats(source):
                label = f"{source} → {target}"
                if source == "pages" or target == "pages":
                    results.append(
                        {
                            "route": label,
                            "status": "manual",
                            "reason": "Apple Pages Automation permission",
                        }
                    )
                    continue
                try:
                    output = convert(sources[source], target)
                    validate(output, target)
                    results.append({"route": label, "status": "passed"})
                except Exception as error:
                    results.append({"route": label, "status": "failed", "error": str(error)})
            print(f"{source}: checked", flush=True)
    report = {
        "total": total,
        "passed": sum(r["status"] == "passed" for r in results),
        "manual": sum(r["status"] == "manual" for r in results),
        "failed": sum(r["status"] == "failed" for r in results),
        "routes": results,
    }
    output = ROOT / ".build/matrix-report.json"
    output.parent.mkdir(exist_ok=True)
    output.write_text(json.dumps(report, indent=2))
    print(json.dumps({k: v for k, v in report.items() if k != "routes"}), flush=True)
    return bool(report["failed"])


if __name__ == "__main__":
    sys.exit(main())
